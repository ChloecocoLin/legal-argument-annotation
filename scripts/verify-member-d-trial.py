"""Offline review checks. Requires jsonschema/openpyxl only for this verification tool."""
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import argparse
parser = argparse.ArgumentParser()
parser.add_argument("--schema-only", action="store_true")
parser.add_argument("--export-only", action="store_true")
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
if not args.export_only:
    from jsonschema import Draft202012Validator
    schema = json.loads((root / "contracts/review/member-d-batches-trial.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    examples = json.loads((root / "contracts/review/member-d-batches-trial.mock.json").read_text(encoding="utf-8"))["examples"]
    for kind, example in examples.items():
        Draft202012Validator({**schema, "$ref": "#/$defs/" + kind}).validate(example)
    print(f"Trial review schemas: {len(examples)} examples valid")
if args.schema_only:
    raise SystemExit(0)
from openpyxl import load_workbook
book = load_workbook(root / "backend/target/member-d-export.xlsx", data_only=False)
assert {"TextVersions", "Elements", "Relations", "Participants", "Provenance", "LongText"} <= set(book.sheetnames)
assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row), "Formula cell found"
rows = list(book["LongText"].values)
chunks = [dict(zip(rows[0], row)) for row in rows[1:]]
parts = sorted((p for p in chunks if p["field"] == "normalizedText"), key=lambda p: int(p["part"]))
assert [int(p["part"]) for p in parts] == list(range(1, len(parts) + 1))
restored = "".join(p["text"] for p in parts)
assert restored == (root / "backend/target/member-d-export-long-text.txt").read_text(encoding="utf-8")
print(f"XLSX: {len(book.sheetnames)} sheets; long text lossless; zero formula cells")
reports = [ET.parse(p).getroot() for p in (root / "backend/target/surefire-reports").glob("TEST-*.xml")]
assert reports and not any(int(r.get("failures", "0")) or int(r.get("errors", "0")) for r in reports)
print(f"Backend: {sum(int(r.get('tests', '0')) for r in reports)} tests; zero failures/errors")
